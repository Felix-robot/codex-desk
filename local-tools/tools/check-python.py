"""Exercise local artifact and numerical libraries without changing user projects."""
import sys, json, tempfile, pathlib, importlib
out = pathlib.Path(sys.argv[1] if len(sys.argv)>1 else tempfile.mkdtemp(prefix='sub2api-tools-'))
out.mkdir(parents=True, exist_ok=True)
modules=['docx','pypdf','pdfplumber','openpyxl','pptx','reportlab','matplotlib','numpy','pandas','scipy','seaborn','shap','umap','nbformat','nbclient','plotly','statsmodels','sklearn','yaml']
for name in modules: importlib.import_module(name)
from docx import Document
from openpyxl import Workbook, load_workbook
from pptx import Presentation
from reportlab.pdfgen import canvas
from pypdf import PdfReader
import pypdfium2 as pdfium
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import numpy as np
from sklearn.linear_model import LinearRegression
import nbformat
from nbclient import NotebookClient
doc=Document(); doc.add_paragraph('DOCX_OK'); doc.save(out/'check.docx'); assert Document(out/'check.docx').paragraphs[0].text=='DOCX_OK'
book=Workbook();book.active['A1']=7;book.active['A2']='=A1*2';book.save(out/'check.xlsx');assert load_workbook(out/'check.xlsx').active['A2'].value=='=A1*2'
ppt=Presentation();slide=ppt.slides.add_slide(ppt.slide_layouts[0]);slide.shapes.title.text='PPTX_OK';ppt.save(out/'check.pptx');assert Presentation(out/'check.pptx').slides[0].shapes.title.text=='PPTX_OK'
c=canvas.Canvas(str(out/'check.pdf'));c.drawString(80,750,'PDF_OK');c.save();assert 'PDF_OK' in PdfReader(out/'check.pdf').pages[0].extract_text()
pdf=pdfium.PdfDocument(out/'check.pdf');pdf[0].render(scale=1).to_pil().save(out/'check-pdf.png');pdf.close()
plt.plot([0,1,2],[1,3,5]);plt.savefig(out/'check-plot.png');plt.close()
model=LinearRegression().fit(np.array([[0],[1],[2]]),np.array([1,3,5]));assert abs(model.predict([[3]])[0]-7)<1e-6
nb=nbformat.v4.new_notebook(cells=[nbformat.v4.new_code_cell("print('NOTEBOOK_OK')")]);nbformat.write(nb,out/'check.ipynb')
# A temporary kernel specification is restricted to this Python environment.
import os
kernel_dir=out/'jupyter'/'kernels'/'desk-python';kernel_dir.mkdir(parents=True,exist_ok=True)
(kernel_dir/'kernel.json').write_text(json.dumps({'argv':[sys.executable,'-m','ipykernel_launcher','-f','{connection_file}'],'display_name':'Desk Python','language':'python'}),encoding='utf-8')
os.environ['JUPYTER_PATH']=str(out/'jupyter')
executed=NotebookClient(nb,timeout=30,kernel_name='desk-python').execute();assert 'NOTEBOOK_OK' in executed.cells[0].outputs[0].text
print(json.dumps({'python':sys.executable,'imported':len(modules),'docx':True,'xlsx':True,'pptx':True,'pdfAndRender':True,'plot':True,'numericalModel':True,'notebookExecution':True,'evidence':str(out)}))
